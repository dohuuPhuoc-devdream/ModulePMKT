from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
import pypdfium2 as pdfium
from PIL import Image,ImageOps,ImageDraw
root=Path('artifacts/phase8')
for path in root.glob('*.pdf'):
    assert path.stat().st_size > 500, f'{path.name} is empty or incomplete'
    reader=PdfReader(path, strict=True)
    assert len(reader.pages)>0
    for page in reader.pages:
        assert float(page.mediabox.width)==842 and float(page.mediabox.height)==595, f'{path.name} has unexpected page size'
    print(path.name,len(reader.pages),'pages valid')
    text='\n'.join(page.extract_text() or '' for page in reader.pages)
    if path.name=='multipage.pdf':
        for value in ('BÁO CÁO KIỂM THỬ TIẾNG VIỆT','ROW0','ROW72','73 dòng'):
            assert value in text, f'multipage.pdf is missing searchable text {value!r}'
    if path.name=='voucher-template.pdf':
        for value in ('PHIẾU NHẬP MUA','DEMO-N-1610','Nhà cung cấp 1','Tổng cộng','Người nhận hàng','Thủ kho'):
            assert value in text, f'voucher-template.pdf is missing searchable text {value!r}'
    doc=pdfium.PdfDocument(str(path))
    for i in range(len(doc)):
        doc[i].render(scale=1.4).to_pil().save(root/(path.stem+f'-{i+1}.png'))
print('QuestPDF report and voucher PDF text is searchable')
for path in root.glob('*.xlsx'):
    wb=load_workbook(path,data_only=False);ws=wb.active
    if path.name!='voucher-template.xlsx':
        assert ws['A1'].value and ws['A2'].value
        print(path.name,ws.max_row-5,'data rows valid')
voucher=load_workbook(root/'voucher-template.xlsx',data_only=False).active
assert voucher['A1'].value=='Đơn vị PicoCorp' and voucher['A4'].value=='PHIẾU NHẬP MUA'
assert voucher['A8'].value=='Nhà cung cấp' and voucher['B8'].value=='Nhà cung cấp 1'
assert voucher.page_setup.orientation=='landscape' and str(voucher.page_setup.paperSize)=='9'
assert voucher.print_area and voucher.max_row>=17
assert voucher['A13'].fill.fgColor.rgb=='FFE8EEF3'
assert any(cell.data_type=='f' and cell.value.startswith('=SUM(') for row in voucher.iter_rows() for cell in row)
print('voucher-template.xlsx designed print layout valid')
ws=load_workbook(root/'multipage.xlsx').active
assert ws.max_row==78 and ws['A78'].value=='ROW72'
assert ws['B6'].value.startswith('=HYPERLINK') and ws['B6'].data_type=='s'
assert abs(ws['D6'].value-0.1234)<0.00000001 and ws['C6'].is_date
ws=load_workbook(root/'inventory.xlsx').active
assert ws['A6'].value=='HANG001' and ws['G6'].value==130
print('INDEPENDENT PDF/XLSX VALIDATION PASSED')
files=list(root.glob('*.png'))
thumbs=[]
for f in files:
    im=Image.open(f).convert('RGB'); im.thumbnail((560,400))
    tile=Image.new('RGB',(580,430),'#dddddd');tile.paste(im,(10,25));ImageDraw.Draw(tile).text((10,5),f.name,fill='black');thumbs.append(tile)
canvas=Image.new('RGB',(580*3,430*((len(thumbs)+2)//3)),'white')
for i,t in enumerate(thumbs):canvas.paste(t,((i%3)*580,(i//3)*430))
canvas.save(root/'contact-sheet.png')
# Cross-check known isolated fixture values independently of the application writer.
for name, expected in [('sales-detail',300),('revenue',300),('sales-analysis',300),('purchase-detail',1000),('purchase-analysis',1000),('receivables',200),('payables',900)]:
    ws=load_workbook(root/(name+'.xlsx')).active
    assert ws.cell(6,ws.max_column).value==expected,(name,expected)
ws=load_workbook(root/'balance.xlsx').active
values={ws.cell(r,1).value:ws.cell(r,4).value for r in range(6,ws.max_row+1)}
assert values['A']==values['SOURCE']==100 and values['CHECK']==0
print('ALL REPORT FIXTURE TOTALS RECONCILED')
